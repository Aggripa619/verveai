"""Builds the email-gated Excel template offered on /blog/demand-forecasting-excel.

Mirrors the post's worked example ("Hoodie A"): moving average, exponential
smoothing, FORECAST.LINEAR, FORECAST.ETS, an Actual-vs-Forecast chart, then
safety stock + reorder point. Output path is the secret download URL that
/api/dl returns after the email gate (see SHEET_URLS there).

Run from verveai/:  python scripts/build-demand-forecast-template.py
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path('public/downloads/f7c2a91e4b3d/Verve-AI-Demand-Forecast-Template.xlsx')

NAVY, TEAL, PURPLE = '132144', '00C9A7', '8048F5'
INPUT_FILL = PatternFill('solid', fgColor='FFF7D6')   # yellow = type here
CALC_FILL = PatternFill('solid', fgColor='F3F4F6')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
FORECAST_FILL = PatternFill('solid', fgColor='E6FAF6')
WHITE_BOLD = Font(bold=True, color='FFFFFF')
THIN = Side(style='thin', color='D1D5DB')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HISTORY = 24          # months of actuals
FUTURE = 6            # months to forecast
FIRST = 6             # first data row on the Forecast sheet
LAST_HIST = FIRST + HISTORY - 1
LAST = LAST_HIST + FUTURE

# Sample "Hoodie A": gentle growth + a Q4 peak, matching the post's 120 -> 200 story.
SAMPLE = [120, 130, 125, 118, 112, 115, 122, 135, 150, 170, 195, 200,
          132, 141, 138, 129, 124, 128, 136, 149, 166, 188, 214, 221]


def month(i: int) -> date:
    y, m = divmod(i, 12)
    return date(2024 + y, m + 1, 1)


def header(ws, row, labels, widths=None):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.fill, c.font, c.border = HEAD_FILL, WHITE_BOLD, BOX
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 32
    for col, w in enumerate(widths or [], start=1):
        ws.column_dimensions[chr(64 + col)].width = w


wb = Workbook()

# ── Start Here ────────────────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Start Here'
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 100
lines = [
    ('Demand Forecasting Template', Font(bold=True, size=20, color=NAVY)),
    ('by Verve AI · getverveai.com/blog/demand-forecasting-excel', Font(italic=True, color='6B7280')),
    ('', None),
    ('How to use it (5 minutes)', Font(bold=True, size=13, color=NAVY)),
    ('1. Go to the Forecast tab and replace the yellow "Units Sold" cells with your last 24 months of sales for ONE product.', None),
    ('   Change the first month in A6 to your own start month; the other months update automatically.', None),
    ('2. Read off the next 6 months of forecasts (green rows) from four methods side by side.', None),
    ('3. Check the "Accuracy (MAPE)" row: the lowest % is the method that fit your history best. Use that one.', None),
    ('4. Go to the Safety Stock tab, enter your supplier lead times, and get your safety stock and reorder point.', None),
    ('', None),
    ('Colour key', Font(bold=True, size=13, color=NAVY)),
    ('Yellow cells = your inputs.  Grey cells = formulas (don\'t overwrite).  Green rows = forecast months.', None),
    ('', None),
    ('The four methods', Font(bold=True, size=13, color=NAVY)),
    ('3-Month Moving Average: average of the previous 3 months. Simple and stable; slow to react to trends.', None),
    ('Exponential Smoothing: weights recent months more. Tune the smoothing factor (alpha) on the Forecast tab: higher = reacts faster.', None),
    ('Linear Trend (FORECAST.LINEAR): fits a straight line through your history. Good for steady growth or decline.', None),
    ('Seasonal ETS (FORECAST.ETS): Excel\'s built-in seasonal model (12-month season). Best when you have a repeating yearly pattern.', None),
    ('   Needs Excel 2016 or later (Windows/Mac/365). In Google Sheets or older Excel this column shows an error; use the other three.', None),
    ('', None),
    ('Forecasting more than a handful of SKUs?', Font(bold=True, size=13, color=NAVY)),
    ('Verve AI connects to Shopify or WooCommerce and forecasts every SKU automatically from your live sales data. Free 14-day trial:', None),
    ('Shopify: https://apps.shopify.com/verve-ai', Font(color=TEAL, underline='single')),
    ('WooCommerce: https://wordpress.org/plugins/verve-ai-inventory-forecasting', Font(color=PURPLE, underline='single')),
]
for i, (text, font) in enumerate(lines, start=2):
    c = ws.cell(row=i, column=2, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if font:
        c.font = font
ws['B24'].hyperlink = 'https://apps.shopify.com/verve-ai?utm_source=excel_template'
ws['B25'].hyperlink = 'https://wordpress.org/plugins/verve-ai-inventory-forecasting'

# ── Forecast ──────────────────────────────────────────────────────────────────
fc = wb.create_sheet('Forecast')
fc['A1'] = 'Demand Forecast: replace the yellow cells with your own monthly sales'
fc['A1'].font = Font(bold=True, size=14, color=NAVY)
fc['A2'] = 'Product / SKU'
fc['B2'] = 'Hoodie A'
fc['A3'] = 'Smoothing factor (alpha, 0.1–0.9)'
fc['B3'] = 0.3
for ref in ('B2', 'B3'):
    fc[ref].fill, fc[ref].border = INPUT_FILL, BOX
fc['A2'].font = fc['A3'].font = Font(bold=True)
fc['B3'].number_format = '0.00'
dv = DataValidation(type='decimal', operator='between', formula1='0.05', formula2='0.95',
                    showErrorMessage=True, error='Enter a value between 0.05 and 0.95')
fc.add_data_validation(dv)
dv.add('B3')

header(fc, FIRST - 1,
       ['Month', 'Units Sold', '3-Month Moving Avg', 'Exponential Smoothing',
        'Linear Trend\n(FORECAST.LINEAR)', 'Seasonal ETS\n(FORECAST.ETS)', 'Period #'],
       widths=[34, 14, 20, 22, 22, 22, 10])

hist_b = f'$B${FIRST}:$B${LAST_HIST}'
# Period numbers, not dates, as the timeline: FORECAST.ETS needs an evenly
# spaced timeline and calendar months aren't (28-31 days).
hist_g = f'$G${FIRST}:$G${LAST_HIST}'
for i in range(HISTORY + FUTURE):
    r = FIRST + i
    future = r > LAST_HIST
    a = fc.cell(row=r, column=1, value=month(0) if r == FIRST else f'=EDATE(A{r - 1},1)')
    a.number_format = 'mmm yyyy'
    b = fc.cell(row=r, column=2, value=None if future else SAMPLE[i])

    # Moving average: forecast for this month = mean of the 3 months before it.
    # Forecast months beyond the first only have the last 3 actuals to go on.
    if r >= FIRST + 3 and not future:
        ma = f'=AVERAGE(B{r - 3}:B{r - 1})'
    elif future:
        ma = f'=AVERAGE($B${LAST_HIST - 2}:$B${LAST_HIST})'
    else:
        ma = None
    fc.cell(row=r, column=3, value=ma)

    # Exponential smoothing: seeded with the first actual; forecasts stay flat
    # at the last smoothed level (that's what simple smoothing predicts).
    if r == FIRST:
        es = f'=B{r}'
    elif not future:
        es = f'=$B$3*B{r - 1}+(1-$B$3)*D{r - 1}'
    elif r == LAST_HIST + 1:
        es = f'=$B$3*B{r - 1}+(1-$B$3)*D{r - 1}'
    else:
        es = f'=D{r - 1}'
    fc.cell(row=r, column=4, value=es)

    fc.cell(row=r, column=5, value=f'=_xlfn.FORECAST.LINEAR(G{r},{hist_b},{hist_g})')
    fc.cell(row=r, column=6, value=f'=_xlfn.FORECAST.ETS(G{r},{hist_b},{hist_g},12)' if future else None)
    g = fc.cell(row=r, column=7, value=i + 1)
    g.font = Font(color='9CA3AF')
    g.alignment = Alignment(horizontal='center')

    for col in range(1, 7):
        c = fc.cell(row=r, column=col)
        c.border = BOX
        if col >= 3:
            c.number_format = '#,##0'
        if future:
            c.fill = FORECAST_FILL
            if col == 1:
                c.font = Font(bold=True)
        elif col == 2:
            c.fill = INPUT_FILL
        elif col >= 3:
            c.fill = CALC_FILL
    if future:
        fc.cell(row=r, column=1).number_format = '"Forecast: "mmm yyyy'

# Accuracy: mean absolute % error of each method over the history it covers.
acc = LAST + 2
fc.cell(row=acc, column=1, value='Accuracy (MAPE, lower = better)').font = Font(bold=True)
spans = {3: FIRST + 3, 4: FIRST + 1, 5: FIRST}
for col, start in spans.items():
    L = chr(64 + col)
    f = (f'=SUMPRODUCT(ABS($B${start}:$B${LAST_HIST}-{L}{start}:{L}{LAST_HIST})/$B${start}:$B${LAST_HIST})'
         f'/ROWS($B${start}:$B${LAST_HIST})')
    c = fc.cell(row=acc, column=col, value=f)
    c.number_format, c.font, c.border = '0.0%', Font(bold=True), BOX
fc.cell(row=acc, column=6, value='n/a (forecast only)').alignment = Alignment(horizontal='center')
fc.cell(row=acc + 1, column=1, value='Best-fitting method').font = Font(bold=True)
best = (f'=INDEX($C${FIRST - 1}:$E${FIRST - 1},MATCH(MIN(C{acc}:E{acc}),C{acc}:E{acc},0))')
fc.cell(row=acc + 1, column=3, value=best).font = Font(bold=True, color=TEAL)
fc.conditional_formatting.add(
    f'C{acc}:E{acc}',
    FormulaRule(formula=[f'C{acc}=MIN($C${acc}:$E${acc})'], fill=PatternFill('solid', fgColor='C7F5EC')),
)
fc.cell(row=acc + 3, column=1,
        value='Tip: if your product has a clear yearly peak (e.g. Q4), compare against the Seasonal ETS column too; '
              'MAPE can\'t score it because it only forecasts future months.').alignment = Alignment(wrap_text=True)
fc.merge_cells(start_row=acc + 3, start_column=1, end_row=acc + 3, end_column=6)
fc.row_dimensions[acc + 3].height = 32
fc.freeze_panes = f'B{FIRST}'

chart = LineChart()
chart.title = 'Actual vs Forecast'
chart.y_axis.title = 'Units'
chart.height, chart.width = 9, 22
chart.add_data(Reference(fc, min_col=2, max_col=6, min_row=FIRST - 1, max_row=LAST), titles_from_data=True)
chart.set_categories(Reference(fc, min_col=1, min_row=FIRST, max_row=LAST))
chart.x_axis.number_format = 'mmm yy'
for s, color in zip(chart.series, [NAVY, '9CA3AF', PURPLE, 'F59E0B', TEAL]):
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = 28000 if color == NAVY else 18000
    s.smooth = False
fc.add_chart(chart, 'H5')

# ── Safety Stock ──────────────────────────────────────────────────────────────
ss = wb.create_sheet('Safety Stock')
ss.column_dimensions['A'].width = 46
ss.column_dimensions['B'].width = 16
ss.column_dimensions['C'].width = 60
ss['A1'] = 'Safety Stock & Reorder Point'
ss['A1'].font = Font(bold=True, size=14, color=NAVY)
ss['A2'] = 'Daily usage comes from the Forecast tab automatically; enter your supplier lead times.'
ss['A2'].font = Font(italic=True, color='6B7280')
rows = [
    ('Average daily usage (units)', f'=AVERAGE(Forecast!B{LAST_HIST - 11}:B{LAST_HIST})/30', 'calc',
     'Last 12 months of actuals ÷ 30'),
    ('Maximum daily usage (units)', f'=MAX(Forecast!B{LAST_HIST - 11}:B{LAST_HIST})/30', 'calc',
     'Busiest month in the last 12 ÷ 30'),
    ('Average lead time (days)', 21, 'input', 'Typical days from placing a PO to stock arriving'),
    ('Maximum lead time (days)', 35, 'input', 'Longest delay you\'ve seen from this supplier'),
    ('', None, None, None),
    ('Safety stock (units)', '=ROUNDUP(B5*B7-B4*B6,0)', 'result',
     '(Max daily usage × Max lead time) − (Avg daily usage × Avg lead time)'),
    ('Reorder point (units)', '=ROUNDUP(B4*B6+B9,0)', 'result',
     '(Avg daily usage × Avg lead time) + Safety stock. Reorder when stock falls to this'),
    ('Next month\'s forecast (best-fitting method)',
     f'=INDEX(Forecast!C{LAST_HIST + 1}:E{LAST_HIST + 1},MATCH(MIN(Forecast!C{acc}:E{acc}),Forecast!C{acc}:E{acc},0))',
     'calc', 'From the Forecast tab: the method with the lowest MAPE'),
]
for i, (label, value, kind, note) in enumerate(rows, start=4):
    if not label:
        continue
    ss.cell(row=i, column=1, value=label).font = Font(bold=kind == 'result')
    c = ss.cell(row=i, column=2, value=value)
    c.border = BOX
    c.number_format = '#,##0.0' if kind == 'calc' else '#,##0'
    c.fill = {'input': INPUT_FILL, 'calc': CALC_FILL, 'result': PatternFill('solid', fgColor='C7F5EC')}[kind]
    if kind == 'result':
        c.font = Font(bold=True, size=12)
    ss.cell(row=i, column=3, value=note).font = Font(color='6B7280')

for sheet in (fc, ss):
    sheet.sheet_view.showGridLines = False
wb.active = 1  # open on the Forecast tab

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f'Wrote {OUT}')
